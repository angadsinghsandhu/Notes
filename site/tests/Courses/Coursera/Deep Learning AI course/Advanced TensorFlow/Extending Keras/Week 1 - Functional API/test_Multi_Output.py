"""Actual UCI workbook parity at selected data boundary; never model execution."""
import ast
import hashlib
from io import BytesIO
import json
from pathlib import Path
import sys
import unittest
from unittest.mock import patch

import nbformat
import numpy as np
import openpyxl
import pandas as pd
import polars as pl

SITE = next(p for p in Path(__file__).resolve().parents if p.name == 'site')
sys.path.insert(0, str(SITE / 'scripts'))
from check_polars import cell_source, read_source

PATH = 'Courses/Coursera/Deep Learning AI course/Advanced TensorFlow/Extending Keras/Week 1 - Functional API/Multi-Output.ipynb'
HASH = '899fa19bc53adb1fdb7a4235cd9c21e791056167135d1de6e3332f2e9a7b90f5'
WORKBOOK = SITE / 'tests/fixtures/real/polars/ENB2012_data.xlsx'
WORKBOOK_HASH = '0089fcffc1415e41e2ff63730cca5280efe54fc43d722dfde1b0aaa808e35dc4'
FEATURES = [f'X{i}' for i in range(1, 9)]


def preprocessing(source):
    tree = ast.parse(cell_source(source, 3).split('# Define model layers.')[0])
    tree.body = [n for n in tree.body if not isinstance(n, (ast.Import, ast.ImportFrom))]
    return compile(tree, '<Excel preprocessing only>', 'exec')


class MultiOutputTest(unittest.TestCase):
    def test_polars_contract_and_raw_preservation(self):
        source, old = read_source(PATH), read_source(PATH, original=True)
        self.assertEqual(hashlib.sha256(old.encode()).hexdigest(), HASH)
        self.assertIn('import polars as pl', cell_source(source, 3))
        self.assertNotIn('import pandas', cell_source(source, 3))
        self.assertIn('polars[excel]', cell_source(source, 3))
        ast.parse(cell_source(source, 3))
        before, after = json.loads(old), json.loads(source)
        nbformat.validate(after)
        self.assertEqual(len(before['cells']), len(after['cells']))
        for i, (a, b) in enumerate(zip(before['cells'], after['cells'])):
            if i == 3:
                self.assertEqual({k: v for k, v in a.items() if k != 'source'}, {k: v for k, v in b.items() if k != 'source'})
                tail = '# Define model layers.'
                self.assertEqual(cell_source(source, 3).split(tail)[1].replace('len(FEATURES)', 'len(train .columns)'), cell_source(old, 3).split(tail)[1])
            else:
                self.assertEqual(a, b)
        self.assertEqual({k: v for k, v in before.items() if k != 'cells'}, {k: v for k, v in after.items() if k != 'cells'})

    def run_preprocessing(self, original=False, payload=None):
        raw = pd.read_excel(WORKBOOK, engine='openpyxl')
        permutation = np.random.RandomState(17).permutation(768)
        order = np.random.RandomState(23).permutation(768)
        seen = []

        def split(rows, **kwargs):
            self.assertEqual(kwargs, {'test_size': 0.2})
            self.assertEqual(len(rows), 768)
            seen.append(rows)
            if isinstance(rows, pd.DataFrame):
                return rows.iloc[order[154:]].copy(), rows.iloc[order[:154]].copy()
            return np.asarray(rows)[order[154:]].tolist(), np.asarray(rows)[order[:154]].tolist()

        def fetch(url, timeout):
            self.assertEqual(url, 'https://archive.ics.uci.edu/ml/machine-learning-databases/00242/ENB2012_data.xlsx')
            self.assertEqual(timeout, 30)
            return BytesIO(WORKBOOK.read_bytes() if payload is None else payload)

        scope = {'pd': pd, 'pl': pl, 'np': np, 'BytesIO': BytesIO, 'urlopen': fetch, 'train_test_split': split}
        with patch.object(pd, 'read_excel', return_value=raw.copy()), \
                patch.object(pd.DataFrame, 'sample', lambda frame, frac: frame.iloc[permutation]), \
                patch.object(np.random, 'permutation', return_value=permutation):
            exec(preprocessing(read_source(PATH, original=original)), scope)
        self.assertEqual(len(seen), 1)
        return scope, raw, permutation, order

    def test_actual_workbook_shuffle_splits_targets_train_only_sample_statistics(self):
        self.assertEqual(hashlib.sha256(WORKBOOK.read_bytes()).hexdigest(), WORKBOOK_HASH)
        book = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
        sheet = book.worksheets[0]
        self.assertEqual((sheet.max_row, sheet.max_column), (1297, 12))
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(sum(any(v is not None for v in row) for row in rows), 769)
        self.assertTrue(all(v is None for row in rows for v in row[10:]))
        book.close()
        old, raw, permutation, order = self.run_preprocessing(True)
        new, _, _, _ = self.run_preprocessing()
        self.assertIsInstance(new['df'], pl.DataFrame, 'workbook still returns Pandas')
        self.assertEqual(new['df'].shape, (768, 10))
        self.assertEqual(new['df'].columns, FEATURES + ['Y1', 'Y2'])
        np.testing.assert_array_equal(new['df'].to_numpy(), old['df'].to_numpy())
        for name, ids in [('train', permutation[order[154:]]), ('test', permutation[order[:154]])]:
            np.testing.assert_array_equal(new[name].to_numpy(), raw.iloc[ids].to_numpy())
            for label, column in zip(new[f'{name}_Y'], ['Y1', 'Y2']):
                self.assertEqual(label.dtype, np.dtype('float64'))
                np.testing.assert_array_equal(label, raw.iloc[ids][column].to_numpy())
            self.assertEqual(new[f'norm_{name}_X'].shape, (len(ids), 8))
            np.testing.assert_allclose(new[f'norm_{name}_X'], old[f'norm_{name}_X'].to_numpy(), rtol=1e-12, atol=1e-12)
        for statistic in ('mean', 'std'):
            self.assertEqual(list(new['train_stats'][statistic]), FEATURES)
            np.testing.assert_allclose(list(new['train_stats'][statistic].values()), old['train_stats'][statistic].to_numpy(), rtol=1e-12)
        self.assertFalse(np.allclose(list(new['train_stats']['mean'].values()), raw[FEATURES].mean().to_numpy()))
        self.assertEqual([str(new['df'][n].dtype) for n in ('X6', 'X8')], ['Int64', 'Int64'])
        self.assertEqual(new['df'].null_count().sum_horizontal().item(), 0)

    def test_bounded_read_and_null_nan_constant_empty_normalization(self):
        source = read_source(PATH)
        self.assertIn('import polars as pl', cell_source(source, 3))
        with self.assertRaisesRegex(ValueError, '2 MiB'):
            self.run_preprocessing(payload=b'x' * (2 * 1024 * 1024 + 1))
        old, raw, _, _ = self.run_preprocessing(True)
        new, _, _, _ = self.run_preprocessing()
        edge = raw.iloc[:5].copy()
        edge.loc[0, 'X1'] = np.nan
        edge.loc[1, 'X2'] = np.nan
        edge['X3'] = 1.0
        frame = pl.DataFrame({n: edge[n].to_numpy() for n in edge.columns}).with_columns(
            pl.when(pl.int_range(pl.len()) == 0).then(None).otherwise(pl.col('X1')).alias('X1'))
        tree = ast.parse(cell_source(source, 3))
        stats = [n for n in tree.body if isinstance(n, ast.Assign) and any(isinstance(t, ast.Name) and t.id == 'train_stats' for t in n.targets)]
        new['train'] = frame
        exec(compile(ast.Module(body=stats, type_ignores=[]), '<Excel stats>', 'exec'), new)
        old['train_stats'] = edge[FEATURES].describe().transpose()
        with np.errstate(invalid='ignore', divide='ignore'):
            np.testing.assert_allclose(new['norm'](frame), old['norm'](edge[FEATURES]).to_numpy(), equal_nan=True)
        self.assertTrue(frame.equals(new['train']))
        self.assertEqual(new['norm'](frame[:0]).shape, (0, 8))


if __name__ == '__main__':
    unittest.main()
